#!/usr/bin/env python3
"""Generate the editable idea brief and the business/evaluation PDF reports."""

from __future__ import annotations

import json
import shutil
import subprocess
import textwrap
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
TODAY = "27 October 2026"
NAVY = "173747"
TEAL = "0D8F89"
INK = "253E4D"
MUTED = "6C808B"
PALE = "EFF5F5"
SAND = "FFF4E2"
SCREEN = ROOT / "assets" / "DrillLoop_Demo_Overview.png"
BENCHMARK = json.loads((ROOT / "demo" / "benchmark_results.json").read_text())

REFS = [
    "[1] FEMA, Homeland Security Exercise and Evaluation Program Doctrine, January 2020: https://www.fema.gov/sites/default/files/2020-04/Homeland-Security-Exercise-and-Evaluation-Program-Doctrine-2020-Revision-2-2-25.pdf",
    "[2] FEMA Preparedness Toolkit, Exercise Evaluation: https://preptoolkit.fema.gov/web/training/evaluation",
    "[3] California Department of Public Health, Power Loss scenario and hospital objectives: https://www.cdph.ca.gov/Programs/EPO/Pages/swmhe_scenarios_power_loss.aspx",
    "[4] Veoci, After-action reporting product page: https://veoci.com/emergency-management/after-action-reporting/",
    "[5] D4H, Emergency response team readiness software: https://www.d4h.com/team-readiness-software",
    "[6] Mid-South Emergency Planning Coalition, 2023 Full Scale Exercise AAR/IP example: https://midsouthepc.org/wp-content/uploads/2024/03/Mid-South-HCC-2023-Non-Hospital-FSE-AAR.pdf",
    "[7] Python release history: https://www.python.org/doc/versions/ ; FastAPI 0.162.0: https://pypi.org/project/fastapi/0.162.0/",
    "[8] React versions: https://react.dev/versions ; PostgreSQL versioning policy: https://www.postgresql.org/support/versioning/",
    "[9] faster-whisper: https://github.com/SYSTRAN/faster-whisper ; MIT English-small model card: https://huggingface.co/Systran/faster-whisper-small.en ; Whisper license: https://github.com/openai/whisper/blob/main/README.md",
    "[10] Qwen3 model licenses and local use: https://qwenlm.github.io/blog/qwen3/ ; Ollama local API: https://github.com/ollama/ollama/blob/main/docs/api.md",
    "[11] pypdf 6.21.2 license metadata (BSD-3-Clause): https://pypi.org/project/pypdf/6.21.2/",
    "[12] PptxGenJS speaker notes: https://gitbrent.github.io/PptxGenJS/docs/speaker-notes.html",
]


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    props.append(shd)


def set_cell_text(cell, value: str, *, bold=False, color=INK, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i], NAVY)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=8)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            shade(cells[i], "F2F7F7" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_text(cells[i], value, size=8)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def setup_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.65)
    sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in (("Title", 34, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 10.5, INK)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    header = sec.header.paragraphs[0]
    header.text = "DRILLLOOP  /  PRODUCT IDEA                                  EUREKA HACKATHON  ·  27 OCT 2026"
    header.style = "Caption"
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    header.runs[0].font.bold = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Synthetic demo  ·  Draft for discussion     |     ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    return doc


def add_callout(doc: Document, heading: str, body: str, fill: str = PALE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(heading + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL if fill == PALE else "A27128")
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def build_idea_docx() -> None:
    doc = setup_doc()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    r = p.add_run("EUREKA  /  PRODUCT CONCEPT")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph("DrillLoop", style="Title")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Evidence-linked exercise reviews and action closure")
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph("A local-first review workspace for hospital preparedness evaluators. The proposed first use case is reviewing records from a hospital tabletop exercise; other sectors remain future candidates.")
    add_callout(doc, "Product thesis", "Turn fragmented exercise evidence into a source-linked timeline, evaluator-reviewed findings, and corrective actions that remain open until an authorized reviewer checks closure evidence.")
    if SCREEN.exists():
        doc.add_picture(str(SCREEN), width=Inches(6.65))
        cap = doc.add_paragraph("DrillLoop hospital demo · interface screenshot · all records are synthetic")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph(f"Prepared {TODAY} · Hackathon MVP concept and implementation", style="Subtitle")
    doc.add_page_break()

    doc.add_heading("Start here: one-minute example", 1)
    doc.add_paragraph("In the fictional hospital exercise, the control-room log records a power failure at 09:02 (EL-01) and essential power restored at 09:06 (EL-02). The exercise target is 60 seconds.")
    add_callout(doc, "What DrillLoop proposes", "The recorded restoration took about four minutes, longer than the one-minute target. Evidence: EL-01 and EL-02 in control_room_log.csv. The cause is not in these records, so it stays unknown. The finding remains proposed until an evaluator reviews it.")
    add_table(doc, ["Source says", "System suggests", "Person decides"], [["Power failed at 09:02; stable at 09:06.", "Possible improvement gap against OBJ-01.", "Accept, edit, or reject the proposal; assign an owner only when known."]], [2.1, 2.2, 2.35])
    doc.add_paragraph("For a glossary, audience-specific reading paths, and a conceptual example for AI agents, see DRILLLOOP_READER_GUIDE.md.")

    doc.add_heading("1. Executive summary", 1)
    doc.add_paragraph("The proposed first use case is a hospital preparedness evaluator reviewing records from a tabletop exercise and preparing an After-Action Report and Improvement Plan (AAR/IP). Notes, event logs, transcripts, and timelines may need reconciliation; timing conflicts need review; and improvement tasks can lack an owner or completion proof. DrillLoop keeps proposed findings connected to source records and leaves decisions with a human evaluator.")
    doc.add_paragraph("The working MVP imports notes, CSV event logs, JSON transcripts, and PDFs with selectable text. It builds a timeline from recorded times, proposes findings for one fictional hospital power-outage fixture, supports evaluator approval or rejection, tracks owners and completion evidence, and exports an AAR/IP PDF. The rules have not been evaluated on an unfamiliar exercise.")
    doc.add_paragraph("The default analysis uses fixed local rules and needs no cloud service. An optional local language model can reword an action draft but cannot create findings or citations. An optional speech-to-text path needs a model installed on the computer. Neither optional model was used in the reported benchmark.")
    doc.add_heading("First use case and roles to validate", 2)
    doc.add_paragraph("Initial segment hypothesis: hospital preparedness teams reviewing tabletop exercises. The likely report preparer is an exercise evaluator or preparedness staff member. A preparedness lead or exercise sponsor is a candidate approver; the budget owner and procurement path are unknown. Municipal emergency teams, airports, civil defense, training providers, defense-support organizations, and critical-infrastructure operators are possible later segments, not validated customers.")
    doc.add_heading("Why now / product value to test", 2)
    doc.add_paragraph("The value hypothesis is that source-level traceability and reviewer-controlled actions can reduce evidence hunting and make follow-up easier to verify. No user interviews, reporting-time baseline, evaluator correction study, pilot interest, or time-savings result is recorded. These measures remain to be established with a pilot partner.")
    doc.add_page_break()

    doc.add_heading("2. Workflow and practical demo scenario", 1)
    add_table(doc, ["Step", "Evaluator workflow", "DrillLoop output"], [
        ["Plan", "Set scenario, date, targets, criteria, roles, teams, and milestones.", "Planning record connects findings to the targets being tested."],
        ["Collect", "Upload notes, event logs, transcripts, and reports.", "Original file, file-change checksum, extracted observations, and source locations."],
        ["Reconstruct", "Put recorded times in order and compare sources.", "Missing or conflicting times stay marked as uncertain."],
        ["Review", "Read the draft and its supporting records; edit, approve, or reject.", "The reviewer’s decision is recorded; proposals remain distinct from approvals."],
        ["Improve", "Assign an owner, due date, and a test for completion.", "The action cannot be approved until its finding is approved and required fields exist."],
        ["Verify", "Check the completion proof before closure.", "A reviewer closes the task only after checking evidence."],
    ], [0.75, 3.0, 2.9])
    doc.add_heading("Fictional hospital prolonged power outage", 2)
    doc.add_paragraph("Exercise HOSP-PO-2026-01 (“Night Shift: Extended Power Loss”) tests four objectives. Included evidence is synthetic and deliberately covers failure, improvement, a positive performance example, conflicts, and missing details.")
    add_table(doc, ["Time / signal", "Synthetic source statement", "Expected treatment"], [
        ["09:02", "Primary utility power supply failed.", "Source-backed event; link to OBJ-01."],
        ["09:06", "Backup generator stabilized and essential circuits restored four minutes later.", "Proposed gap against the 60-second objective."],
        ["09:03 → 09:10", "Radio acknowledgement request; ICU acknowledgement arrived by runner after seven minutes.", "Proposed communications finding; evaluator also reports 09:11, unresolved."],
        ["09:02 → 09:08", "Outage start and facilities arrival; incident log says response time was two minutes.", "Show the discrepancy; retain the original source value."],
        ["09:18", "ICU team completed cold-chain relocation and logged chain of custody.", "Proposed strength against OBJ-03."],
        ["Debrief", "Radio follow-up action has no owner; separate note says generator lagged, without time or measure.", "Owner remains blank; ambiguous note is not promoted."],
    ], [0.85, 3.45, 2.35])
    doc.add_paragraph("The provided transcript includes a synthetic 0.58-confidence ‘two’ → ‘too’ recognition substitution. It stays uncertain and the incorrect-response-time draft cites the independent event log, not that transcript phrase.")
    doc.add_page_break()

    doc.add_heading("3. AI and evidence architecture", 1)
    doc.add_paragraph("The workflow is a fixed local pipeline: parse source → extract observations → detect candidate findings → map to a stated objective → attach exact observation IDs → surface timing/recognition uncertainty → draft an action → wait for a human decision. The same evidence IDs resolve to file name and row, line, page, or transcript segment.")
    add_callout(doc, "Evidence rule", "Every displayed candidate includes supporting evidence. If the source is incomplete, the gap is explicit. A model may not fill it with a plausible event or resolve a time conflict.", SAND)
    doc.add_heading("MVP technical choices", 2)
    add_table(doc, ["Layer", "Selection", "Rationale / boundary"], [
        ["Server and data checks", "FastAPI 0.162.0 · Pydantic 2.14.1", "Receives website requests and checks that uploaded data has the expected fields."],
        ["Saved data", "SQLAlchemy 2.0.48 · SQLite default · PostgreSQL 18.6 Compose", "One data-access path supports a local file database or a deployed server database."],
        ["Website screen", "React 19.3.1 · TypeScript 5.9.3 · Vite 8.3.5", "Pages people use to plan exercises, review evidence, and manage tasks."],
        ["PDF import", "pypdf 6.21.2 (BSD-3-Clause)", "Selectable text only; chosen over an AGPL/commercial-dual-license parser."],
        ["Speech-to-text (ASR)", "faster-whisper 1.2.1 optional (MIT)", "Runs locally when a model is installed; it creates a draft transcript and does not identify speakers."],
        ["Optional writing model (LLM)", "Ollama + Qwen3-4B optional (Apache 2.0 model weights)", "May reword an action draft only; not used in the recorded evaluation."],
        ["Reports", "python-docx 1.2.0 · ReportLab 5.1.0 · PptxGenJS 4.0.1", "Editable DOCX/PPTX; PDF report/export and speaker notes."],
        ["Quality", "pytest 9.1.1 · Ruff 0.17.2", "Parser, status transition, citation, and API tests."],
    ], [1.25, 2.65, 2.75])
    doc.add_heading("Why fixed workflow instead of agents and retrieval", 2)
    doc.add_paragraph("This MVP uses a short list of pre-written rules. The same input gives the same draft, so a reviewer can inspect why it appeared. The sample has only a few files, so a separate search system for similar text would add complexity without a clear benefit. A tool such as LangGraph could be considered later if teams need long-running review queues or several people must sign off in sequence.")
    doc.add_page_break()

    doc.add_heading("4. Differentiation and competitive context", 1)
    doc.add_paragraph("After-action reports and corrective-action workflows are available in existing products. Veoci publicly describes after-action reporting and follow-up; D4H serves emergency-response readiness and related team workflows. FEMA’s Homeland Security Exercise and Evaluation Program (HSEEP) provides exercise-planning and evaluation guidance. A report template alone is not a unique product claim. No same-workflow comparison with Veoci or another alternative has been run.")
    add_table(doc, ["Option", "Publicly visible scope", "DrillLoop implication"], [
        ["FEMA HSEEP / Preparedness Toolkit", "Federal exercise-planning and evaluation guidance, including after-action and improvement planning.", "Use familiar concepts; do not imply official status or replace an organization’s required reports."],
        ["Veoci", "Emergency-management platform; after-action reporting, timestamping, records, and corrective-action follow-up.", "Direct competitor and integration/market-fit question; evidence trace must be tested in product demos."],
        ["D4H", "Emergency response team readiness and adjacent incident/personnel/equipment workflows.", "Adjacent incumbent; confirm evaluator evidence-review and closure workflows with customers."],
        ["DrillLoop positioning hypothesis", "Exact source observation → proposed finding → reviewer decision → closure evidence.", "Potential niche: evidence-to-review trace, uncertainty, and action closure. Uniqueness is unvalidated."],
    ], [1.65, 2.75, 2.25])
    doc.add_heading("Differentiator to validate", 2)
    doc.add_paragraph("For each finding, the evaluator can open the exact source text at its filename and row/line/page/segment, see the claim separated from interpretation, confirm or reject it, and later inspect reviewer-approved closure evidence. The differentiator hypothesis is trustworthy traceability across the review lifecycle—not automated prose by itself.")
    doc.add_paragraph("Research references include a public hospital power-loss exercise scenario from California Department of Public Health and a public health-coalition exercise AAR/IP example. These support scenario familiarity and report format; they do not validate product demand.")
    doc.add_page_break()

    doc.add_heading("5. Feasibility, security, and governance", 1)
    doc.add_heading("Feasibility", 2)
    doc.add_paragraph("The app runs without a cloud AI service. It uses SQLite for the local demo and can use PostgreSQL in Docker Compose. The website lets a user create an exercise, upload evidence, inspect proposed findings, review actions, and download a PDF report. The server checks uploads, runs the fixed analysis rules, and records reviewer decisions. Optional speech-to-text and local text-generation paths were not used in the benchmark.")
    doc.add_heading("Security boundary", 2)
    for item in [
        "The demo uses one shared reviewer password. It has no named-user login or role-specific permissions.",
        "Original files and extracted text are saved in the configured database. The demo has no configured encryption at rest, file scanning, retention/deletion controls, tenant isolation, or security assessment.",
        "The packaged demo and benchmark use fictional records. The system has not been validated for protected patient or operational data. A deployment handling such records would require approved hosting, encryption and key management, role controls, access auditing, retention/deletion rules, dependency scanning, and threat modeling.",
        "The default workflow makes no cloud model call. Optional faster-whisper runs with a locally staged model; optional Ollama is configured as a local/private service for action wording. Neither optional path was included in the benchmark.",
    ]:
        bullet(doc, item)
    doc.add_heading("Human review and professional judgment", 2)
    doc.add_paragraph("An authorized reviewer must approve or reject a proposed finding. An action needs an approved finding, owner, and due date before action approval. Closing it also requires reviewer confirmation and completion evidence. The report borrows concepts from FEMA’s Homeland Security Exercise and Evaluation Program (HSEEP), but it is not FEMA-certified.")
    doc.add_page_break()

    doc.add_heading("6. Evaluation, commercialization, and roadmap", 1)
    doc.add_heading("Recorded technical benchmark", 2)
    add_table(doc, ["Measure", "Observed", "Interpretation"], [
        ["Observation extraction", "8/8 · precision 1.00 · recall 1.00", "Labeled CSV rows in one synthetic fixture."],
        ["Finding detection", "5/5 · precision 1.00 · recall 1.00", "Closed ground-truth finding set; not generalization."],
        ["Exact cited event IDs", f"{BENCHMARK['citations_correct']}/{BENCHMARK['citations_expected']}", "All labeled supporting IDs were present."],
        ["Timeline row times", "8/8 exact matches", "Event-log rows only; conflicts remain surfaced separately."],
        ["Action register fields", f"{BENCHMARK['corrective_action_completeness']['complete_required_fields']}/{BENCHMARK['corrective_action_completeness']['possible_required_fields']}", "One intentionally missing owner."],
        ["Rule route latency", f"p50 {BENCHMARK['latency_ms']['p50']} ms · p95 {BENCHMARK['latency_ms']['p95']} ms", "Local TestClient call; no LLM or ASR inference."],
    ], [1.55, 1.7, 3.4])
    doc.add_paragraph("Manual evaluator performance, report-preparation time, correction workload, performance on an unfamiliar exercise, ASR word error rate, local LLM quality, and monetary savings have not been measured. The results are a technical check of the synthetic fixture, not field-accuracy evidence.")
    doc.add_heading("Commercial hypotheses", 2)
    doc.add_paragraph("The first commercial model proposed for discovery is a per-exercise pilot fee. No price, delivery cost, staffing plan, customer volume, or willingness-to-pay evidence is established. Annual subscriptions, private deployment, custom templates, and support services remain options to evaluate after user and buyer discovery.")
    doc.add_paragraph("Time saved and monetary value are unknown. A future business case would need a customer-measured preparation-time baseline, verified change in staff effort, finance-approved labor costs, and evidence that capacity can be redeployed or cash costs reduced.")
    doc.add_heading("Roadmap", 2)
    add_table(doc, ["Stage", "Proposed work", "Evidence required to proceed"], [
        ["Discovery", "Interview 2–3 prospective hospital preparedness users and observe one current review process with permission.", "Role map, baseline measures, data permissions, and a willing pilot contact."],
        ["Paired pilot", "Compare the current process with DrillLoop on the same permitted exercise records; include a tabletop exercise and uncertainty cases.", "Partner-defined thresholds for preparation time, reviewer corrections, citations, and action completeness."],
        ["Generalization and deployment review", "Evaluate evaluator-labeled evidence from an unfamiliar exercise; compare missed/incorrect findings and review effort. Review data flow and access controls.", "Independent labels, false-positive/missed-finding results, security and retention requirements."],
    ], [1.0, 3.0, 2.65])
    doc.add_paragraph("The Eureka form identifies the lead function as Product–Synergy / AI Engineering. Specific delivery allocation, emergency-preparedness experience, and an adviser or pilot partner are not documented. The proposed stages have no confirmed schedule or staffing commitment.")
    doc.add_heading("Key risks", 2)
    doc.add_paragraph("Incomplete inputs; timestamp ambiguity; ASR errors; rules that overfit; confidential evidence exposure; weak adoption; existing platforms that already satisfy the need; and the risk that buyers pay only for convenience instead of measurable improvement. Pilot success criteria should emphasize source-supported findings, reviewer workload, action ownership, and independently verifiable closure.")
    doc.add_page_break()

    doc.add_heading("7. Sources and review notes", 1)
    doc.add_paragraph("The public references below informed workflow and positioning. Accessed or version-checked 26 October 2026. Competitor features are summarized from public product materials, not a scored product teardown. A potential fit with Presight public-safety work is a hypothesis; no integration or customer demand is established.")
    for ref in REFS:
        p = doc.add_paragraph(ref)
        p.paragraph_format.space_after = Pt(7)
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(MUTED)
    add_callout(doc, "Review disclaimer", "DrillLoop provides a working, synthetic demonstration. It does not make safety-critical decisions or replace exercise evaluator judgment, organizational procedures, or jurisdictional reporting requirements.", SAND)
    doc.core_properties.title = "DrillLoop Idea Document"
    doc.core_properties.subject = "Emergency exercise evaluation and corrective-action concept"
    doc.core_properties.author = "DrillLoop project team"
    doc.save(ROOT / "DrillLoop_Idea_Document.docx")


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Brand", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=colors.HexColor("#0D8F89"), spaceAfter=8))
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=29, leading=34, alignment=TA_LEFT, textColor=colors.HexColor("#173747"), spaceAfter=9))
    styles.add(ParagraphStyle(name="DeckSubtitle", parent=styles["Normal"], fontSize=13, leading=18, textColor=colors.HexColor("#6C808B"), spaceAfter=12))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#173747"), spaceBefore=8, spaceAfter=8, keepWithNext=True))
    styles.add(ParagraphStyle(name="Subsection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0D8F89"), spaceBefore=7, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name="BodyDL", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=13, textColor=colors.HexColor("#253E4D"), spaceAfter=6))
    styles.add(ParagraphStyle(name="SmallDL", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.6, leading=10.5, textColor=colors.HexColor("#6C808B"), spaceAfter=5))
    styles.add(ParagraphStyle(name="ReferenceDL", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.3, leading=7.6, textColor=colors.HexColor("#6C808B"), spaceAfter=2.2))
    styles.add(ParagraphStyle(name="CalloutDL", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=13, textColor=colors.HexColor("#173747"), backColor=colors.HexColor("#EFF5F5"), borderPadding=8, spaceBefore=5, spaceAfter=10))
    styles.add(ParagraphStyle(name="CoverMeta", parent=styles["Normal"], fontSize=8, leading=11, textColor=colors.HexColor("#6C808B"), spaceBefore=7))
    return styles


def pdf_table(headers: list[str], rows: list[list[str]], widths: list[float], styles, font=7.6):
    data = [[Paragraph(f"<b>{escape(item)}</b>", styles["SmallDL"]) for item in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(value)).replace("\n", "<br/>"), styles["SmallDL"]) for value in row])
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#173747")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7F7")]),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d8e3e5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return t


def pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#d9e4e5"))
    canvas.line(.62 * inch, .48 * inch, 7.88 * inch, .48 * inch)
    canvas.setFont("Helvetica", 7.2)
    canvas.setFillColor(colors.HexColor("#758892"))
    canvas.drawString(.62 * inch, .3 * inch, "DRILLLOOP  ·  Synthetic exercise  ·  For discussion")
    canvas.drawRightString(7.88 * inch, .3 * inch, f"Page {doc.page}")
    canvas.restoreState()


def cover(story, styles, title, subtitle, meta):
    story += [Spacer(1, .35 * inch), Paragraph("DRILLLOOP  /  EUREKA HACKATHON", styles["Brand"]),
              Paragraph(title, styles["CoverTitle"]), Paragraph(subtitle, styles["DeckSubtitle"]),
              Paragraph(meta, styles["CoverMeta"]), Spacer(1, .15 * inch)]
    if SCREEN.exists():
        img = Image(str(SCREEN), width=6.7 * inch, height=4.61 * inch)
        img.hAlign = "CENTER"
        story += [img, Paragraph("Product interface · screenshot from the working synthetic demo", styles["SmallDL"])]
    story += [Spacer(1, .15 * inch), Paragraph("Fictional data · evidence-linked drafts · human reviewer decisions", styles["CalloutDL"]), PageBreak()]


def build_idea_pdf() -> None:
    styles = pdf_styles()
    story = []
    cover(story, styles, "Idea document", "Evidence-linked exercise reviews and action closure", f"Prepared {TODAY} · Product concept and working hackathon MVP")
    story += [Paragraph("Start here: a plain-language example", styles["Section"]), Paragraph("The fictional hospital log says power failed at 09:02 (EL-01) and essential circuits were restored at 09:06 (EL-02). The exercise target is 60 seconds. DrillLoop can propose that restoration took about four minutes and cite both records. It cannot tell why the delay happened from these records. The evaluator decides whether the proposal is fair.", styles["BodyDL"]), Paragraph("An After-Action Report and Improvement Plan (AAR/IP) says what happened and what the team will improve. The project borrows selected ideas from FEMA’s Homeland Security Exercise and Evaluation Program (HSEEP); it is not FEMA-certified.", styles["BodyDL"]), Paragraph("For a glossary and an example of safe AI-agent output, see DRILLLOOP_READER_GUIDE.md.", styles["CalloutDL"])]
    story += [Paragraph("Executive summary", styles["Section"]), Paragraph("The proposed first use case is a hospital preparedness evaluator reviewing records from a tabletop exercise and preparing an After-Action Report and Improvement Plan (AAR/IP). DrillLoop is a local-first review workspace that ties each draft finding to exact source evidence, maps it to a pre-stated objective, makes uncertainty visible, and keeps closure reviewer-controlled. The current deterministic rules were built around one fictional hospital outage fixture; behavior on an unfamiliar exercise has not been evaluated.", styles["BodyDL"]), Paragraph("Core user and workflow", styles["Subsection"]), Paragraph("Exercise lead configures scenario, date, duration, objectives, criteria, evaluator roles, participants, and milestones. Evaluators upload evidence. The system normalizes explicit times, surfaces conflicts, proposes evidence-backed findings/actions, and prepares a reviewer copy. Authorized reviewers edit/approve/reject and verify closure evidence.", styles["BodyDL"]), Paragraph("The packaged MVP supports TXT/MD, CSV, JSON transcript/event records, selectable-text PDF, optional local audio transcription, exercise setup, timeline, source download, finding review, corrective-action updates, and PDF AAR/IP export. SQLite is the local default; PostgreSQL is available in Compose.", styles["BodyDL"]), Paragraph("Safety rule", styles["CalloutDL"]), Paragraph("Candidate findings remain proposed until a human decision. The workflow does not invent an event, choose between inconsistent times, or close an action without closure evidence. HSEEP concepts inform report structure; DrillLoop is not FEMA-certified.", styles["BodyDL"]),
             Paragraph("First use case and roles to validate", styles["Section"]), Paragraph("Initial segment hypothesis: hospital preparedness teams reviewing tabletop exercises. The likely report preparer is an exercise evaluator or preparedness staff member. A preparedness lead or exercise sponsor is a candidate approver; the budget owner and procurement path are unknown. Municipal teams, airports, civil defense, training providers, defense-support organizations, and critical-infrastructure operators are possible later segments, not validated customers.", styles["BodyDL"]), Paragraph("The value hypothesis is less evidence hunting and more reviewable follow-up. No user interviews, reporting-time baseline, evaluator correction study, pilot interest, or time-savings result is recorded.", styles["BodyDL"]),
             Paragraph("Workflow and synthetic exercise", styles["Section"]), Paragraph("The fictional hospital outage exercise (HOSP-PO-2026-01) tests emergency power transfer within 60 seconds, radio acknowledgement within three minutes, patient-care continuity, and timestamp reconciliation. Synthetic source records contain a 09:02 power failure, essential power restored at 09:06, radio acknowledgement delay, a two-minute reported response against 09:02/09:08 source times, a completed cold-chain procedure, and an unassigned action owner.", styles["BodyDL"])]
    scenario_rows = [
        ["Signal", "Expected treatment"],
        ["Backup generator stable at 09:06 after 09:02 failure", "Improvement draft linked to OBJ-01 and CSV row IDs."],
        ["09:03 request; 09:10 log and 09:11 watch note", "Communication draft; conflict surfaced, not resolved."],
        ["Incident log: two minutes; outage 09:02, arrival 09:08", "Discrepancy draft; cite all relevant rows."],
        ["Cold-chain procedure completed at 09:18", "Proposed strength linked to OBJ-03."],
        ["Action owner blank; vague untimed ‘generator lagged’ note", "Keep owner unassigned and vague statement unpromoted."],
        ["ASR segment says ‘too minutes’ at 0.58 confidence", "Preserve uncertainty; rely on separate source log for time discrepancy."],
    ]
    story.append(pdf_table(scenario_rows[0], scenario_rows[1:], [2.55 * inch, 4.15 * inch], styles))
    story += [Paragraph("AI architecture and evidence controls", styles["Section"]), Paragraph("Parsers create exact observation records, then a fixed local rules pipeline detects candidate findings, maps them to exercise objectives, attaches specific observation IDs, flags time/confidence issues, and drafts corrective actions. The UI resolves citations to the original file and row, line, page, or transcript segment. The included rules are tied to the synthetic fixture; performance on unfamiliar exercises is untested. This small fixture does not need a vector database. LangGraph is omitted because the review stages are fixed and explicit.", styles["BodyDL"]), Paragraph("Optional local Ollama/Qwen3 may rewrite corrective-action wording only. It cannot create candidate events, change evidence IDs, or approve findings. The optional faster-whisper route uses a staged local model and CPU int8. The default workflow makes no cloud call; model quality and ASR accuracy were not tested in the included benchmark.", styles["BodyDL"]), Paragraph("Technology and deployment", styles["Section"]), Paragraph("Pinned stack reviewed 26 October 2026: Python 3.12.16 target; FastAPI 0.162.0; Pydantic 2.14.1; SQLAlchemy 2.0.48; PostgreSQL 18.6; React 19.3.1; TypeScript 5.9.3; Vite 8.3.5; pypdf 6.21.2 (BSD-3-Clause); pytest 9.1.1; Ruff 0.17.2; python-docx 1.2.0; ReportLab 5.1.0; PptxGenJS 4.0.1. The local runtime was Python 3.12.13. Exact compatibility and licenses are in the technical design.", styles["BodyDL"]), Paragraph("Security, privacy, and governance", styles["Section"]), Paragraph("The MVP stores uploaded source bytes, extracted text, and hashes in its configured database. It uses a shared reviewer key, not production identity. It lacks encryption-at-rest configuration, retention/deletion controls, malware scanning, SSO/RBAC, tenant isolation, and a security assessment. The packaged demo uses fictional records and has not been validated for protected patient or operational data. The default workflow makes no cloud model call; optional ASR is local and Ollama is configured as a local/private service. A real deployment needs approved hosting, organization identity, least-privilege review roles, encryption, backup/retention/deletion controls, access auditing, dependency scanning, and security review.", styles["BodyDL"]), PageBreak()]
    story += [Paragraph("Differentiation and competitive context", styles["Section"]), Paragraph("FEMA HSEEP establishes exercise objectives, evaluation, AAR/IP, and improvement-planning concepts. Veoci publicly promotes AAR generation, timestamps, communications transcripts, and corrective-action follow-up; D4H serves emergency response readiness and adjacent incident/personnel/equipment workflows. Report generation and action tracking alone are not unique. DrillLoop's hypothesis is an end-to-end trace from exact source observation to human decision to verified closure. This must be validated with product demos and customer workflow interviews.", styles["BodyDL"]), Paragraph("Evaluation and current results", styles["Section"]), Paragraph(f"On one closed fictional fixture: observation extraction {BENCHMARK['observations']['true_positive']}/{BENCHMARK['observations']['expected']}; finding detection {BENCHMARK['findings']['true_positive']}/{BENCHMARK['findings']['expected']}; exact event-row citations {BENCHMARK['citations_correct']}/{BENCHMARK['citations_expected']}; timeline timestamps {BENCHMARK['timeline']['timestamp_matches']}/{BENCHMARK['timeline']['timestamp_rows']}; action checklist fields {BENCHMARK['corrective_action_completeness']['complete_required_fields']}/{BENCHMARK['corrective_action_completeness']['possible_required_fields']}. This does not establish field quality or human time saved.", styles["BodyDL"]), Paragraph("Manual evaluation was not timed; no real evaluators participated. Local LLM and ASR quality/latency were not run. A next pilot should measure extraction precision/recall, semantic citation support, false/unsupported claims, timeline agreement, reviewer correction workload, report-preparation time, action completeness, closure verification, and resource/latency on consented synthetic or approved exercise records.", styles["BodyDL"]), Paragraph("Commercial model and assumptions", styles["Section"]), Paragraph("The proposed first commercial model to investigate is a per-exercise pilot fee. No price, customer volume, delivery cost, staffing plan, or willingness-to-pay evidence is established. Annual organization subscriptions, private deployment, custom templates, and support are later options to evaluate. A private deployment would require customer-specific security, integration, and support review.", styles["BodyDL"]), Paragraph("Time saved and monetary value are unknown. Measure current preparation and review effort with a prospective partner before estimating impact. Any future labor-cost calculation should use finance-approved inputs and distinguish staff capacity from cash savings.", styles["BodyDL"]), Paragraph("Roadmap", styles["Section"]), Paragraph("Proposed discovery: interview 2–3 prospective hospital-preparedness users and observe a current review with permission. Proposed pilot: compare current practice and DrillLoop on the same permitted records, then evaluate evaluator-labeled evidence from an unfamiliar exercise. Later work on roles, retention, integrations, or optional AI depends on partner requirements and evidence. No partner, schedule, staffing commitment, or success threshold is confirmed.", styles["BodyDL"]), Paragraph("Research and sources", styles["Section"])]
    for ref in REFS:
        story.append(Paragraph(escape(ref), styles["ReferenceDL"]))
    story.append(Paragraph("Synthetic demonstration only. Professional evaluator judgment and organizational procedures remain authoritative.", styles["CalloutDL"]))
    doc = SimpleDocTemplate(str(ROOT / "DrillLoop_Idea_Document.pdf"), pagesize=letter, rightMargin=.62*inch, leftMargin=.62*inch, topMargin=.65*inch, bottomMargin=.65*inch, title="DrillLoop Idea Document")
    doc.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)


def build_business_pdf() -> None:
    styles = pdf_styles()
    story = [
        Paragraph("DRILLLOOP  /  COMMERCIAL HYPOTHESES", styles["Brand"]),
        Paragraph("Business case", styles["CoverTitle"]),
        Paragraph("Customer value and commercial validation plan", styles["DeckSubtitle"]),
        Paragraph(f"Prepared {TODAY} · No prices, costs, sales volumes, or forecast are established.", styles["CoverMeta"]),
        Spacer(1, 8),
        Paragraph("What this document says", styles["Section"]),
        Paragraph("DrillLoop is a proposed local-first workspace. The initial segment hypothesis is hospital preparedness teams reviewing tabletop exercises. The likely user is an exercise evaluator or preparedness staff member; a preparedness lead or exercise sponsor is a candidate report approver. The budget owner and procurement path are unknown. Other emergency-management segments are possible future candidates, not validated customers.", styles["BodyDL"]),
        Paragraph("What the product might add", styles["Section"]),
        Paragraph("Established platforms already offer after-action reports and corrective-action tracking. DrillLoop will test whether hospital evaluators can open the exact source record behind a proposed finding, see uncertainty, record a reviewer decision, and later check completion evidence. This is a hypothesis to validate, not a claim that those features are unique. A possible fit with Presight public-safety work needs internal validation; no integration or customer demand is established.", styles["BodyDL"]),
        Paragraph("First commercial model to investigate: a per-exercise pilot fee. No price, delivery cost, staffing plan, customer volume, or willingness-to-pay evidence is established.", styles["BodyDL"]),
        Paragraph("Commercial options to validate", styles["Section"]),
    ]
    story.append(pdf_table(["Model", "Current status", "Questions to answer"], [
        ["Per-exercise pilot", "First model to investigate; no price set", "Who buys, what procurement route applies, and what one exercise requires."],
        ["Annual organization subscription", "Later hypothesis; not priced", "Exercise volume, reviewer roles, support, and renewal needs."],
        ["Private deployment", "Later hypothesis; not priced", "Hosting, identity, security, integration, and ongoing support requirements."],
        ["Custom templates or services", "Later hypothesis; not priced", "Configuration effort, domain expertise, facilitation, and training needs."],
    ], [1.6*inch, 1.7*inch, 3.4*inch], styles))
    story += [
        PageBreak(),
        Paragraph("Inputs required before modeling economics", styles["Section"]),
        Paragraph("There are no confirmed prices, purchase volumes, delivery capacity, or operating costs, so this document does not present a revenue forecast. Gather these inputs with prospective users, buyers, and finance before preparing a scenario.", styles["BodyDL"]),
    ]
    story.append(pdf_table(["Input", "Evidence to collect"], [
        ["Buyer and purchase", "Identify report preparer, approver, budget owner, purchasing route, and actual willingness to pay."],
        ["Exercise volume", "Use an organization-provided baseline; do not assume annual exercise counts."],
        ["Delivery effort", "Record setup, template configuration, evaluator support, security review, and integration work."],
        ["Operating cost", "Use finance-approved staffing, hosting, support, legal, and security inputs."],
        ["Customer value", "Measure preparation time, corrections, citation quality, and action follow-up; distinguish capacity from cash savings."],
    ], [1.45*inch, 5.25*inch], styles))
    story += [
        Paragraph("Proposed validation plan", styles["Section"]),
    ]
    story.append(pdf_table(["Question", "Proposed evidence"], [
        ["Who owns the workflow and budget?", "Interview prospective hospital-preparedness users; identify report preparer, approver, and budget owner."],
        ["Does traceability improve the current process?", "Compare current practice and DrillLoop on the same permitted exercise records; measure preparation time, evidence lookup, and reviewer corrections."],
        ["How does it compare with an alternative?", "Select one relevant alternative with a prospective partner and compare source-location precision, reviewer decisions, and closure evidence."],
        ["Does it generalize?", "Use evaluator-labeled records from an unfamiliar exercise; report missed and incorrect findings, unsupported claims, and timeline agreement."],
        ["Can it be deployed safely?", "Confirm data rules, identity, retention, access logging, hosting, and security requirements before using non-fictional records."],
    ], [2.05*inch, 4.65*inch], styles))
    story += [
        Paragraph("Known limits and risks", styles["Section"]),
        Paragraph("No customer interviews, baseline, buyer, partner, price, operating-cost model, or willingness-to-pay evidence is documented. Existing tools may already meet the need. The current rules were written for one fictional exercise and may not work on new scenarios. Evidence may be sensitive, and evaluators may not trust AI suggestions. A future study should measure source support, reviewer edits, action ownership, independent closure checks, and actual operating effort.", styles["BodyDL"]),
        Paragraph("Research context", styles["Section"]),
    ]
    for ref in REFS[:6]:
        story.append(Paragraph(escape(ref), styles["ReferenceDL"]))
    doc = SimpleDocTemplate(str(ROOT / "DrillLoop_Business_Case.pdf"), pagesize=letter, rightMargin=.62*inch, leftMargin=.62*inch, topMargin=.62*inch, bottomMargin=.65*inch, title="DrillLoop Business Case")
    doc.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)

def build_evaluation_pdf() -> None:
    styles = pdf_styles()
    observations = BENCHMARK["observations"]
    findings = BENCHMARK["findings"]
    actions = BENCHMARK["corrective_action_completeness"]
    timeline = BENCHMARK["timeline"]
    latency = BENCHMARK["latency_ms"]
    story = [
        Paragraph("DRILLLOOP  /  TECHNICAL EVALUATION", styles["Brand"]),
        Paragraph("Evaluation report", styles["CoverTitle"]),
        Paragraph("What the synthetic test checked — and what it did not", styles["DeckSubtitle"]),
        Paragraph(f"Run date {TODAY} · Fixture: {BENCHMARK['fixture_label']}", styles["CoverMeta"]),
        Spacer(1, 8),
        Paragraph("Short answer", styles["Section"]),
        Paragraph("The rules matched the labels in the one fictional hospital exercise they were written for: 8 of 8 event rows, 5 of 5 expected findings, and 9 of 9 expected evidence-record IDs. This shows that the demo pipeline works on its sample. It does not show how well the product works on a new exercise or whether a professional evaluator agrees with each conclusion.", styles["BodyDL"]),
        Paragraph("How to read the scores", styles["Section"]),
        Paragraph("Precision asks: of the items the system proposed, what share matched the expected labels? Recall asks: of the items in the expected labels, what share did the system find? Here, all five expected finding codes were proposed and no extra code was produced, so both scores are 1.00 on this fixture. The labels were built for this same scenario; perfect fixture scores do not mean perfect real-world accuracy.", styles["BodyDL"]),
        Paragraph("A citation match checks that the expected record ID is linked. It does not prove that the record truly supports the sentence. That semantic question still needs a human reviewer.", styles["BodyDL"]),
        Paragraph("Worked example: power restoration", styles["Subsection"]),
        Paragraph("The log says power failed at 09:02 (EL-01) and essential circuits were restored at 09:06 (EL-02). The exercise target is 60 seconds. Those records support a proposed finding that restoration took about four minutes and missed the target. They do not establish the cause. The benchmark checks that the proposal and record IDs are present; an evaluator must judge whether the interpretation is fair.", styles["BodyDL"]),
        Paragraph("What the fixture contains", styles["Section"]),
        Paragraph("Three invented files: an eight-row event log, evaluator notes with one untimed vague statement and a 09:10/09:11 disagreement, and a pre-generated transcript with an intentionally mistaken ‘too minutes’ phrase at 0.58 confidence. The transcript is not a result of audio transcription. The written labels are in demo/ground_truth.json.", styles["BodyDL"]),
        Paragraph("Recorded technical results", styles["Section"]),
    ]
    rows = [
        ["Check", "Result", "Plain-language meaning"],
        ["Event-row extraction", f"{observations['true_positive']}/{observations['expected']} · precision {observations['precision']:.2f} · recall {observations['recall']:.2f}", "All eight labeled log rows were found in this sample."],
        ["Finding detection", f"{findings['true_positive']}/{findings['expected']} · precision {findings['precision']:.2f} · recall {findings['recall']:.2f}", "All five expected finding labels were produced; no extra label appeared."],
        ["Unsupported proposals", f"{BENCHMARK['unsupported_finding_count']}/{findings['predicted']}", "Closed-fixture code/evidence check only; no human judged whether the claim was fair."],
        ["Proposals with no source link", f"{BENCHMARK['uncited_finding_count']}/{findings['predicted']}", "All five proposals included at least one record ID."],
        ["Evidence-record links", f"{BENCHMARK['citations_correct']}/{BENCHMARK['citations_expected']}", "Expected IDs were linked. This is not a semantic support rating."],
        ["Unresolved / unplaced times", f"{timeline['unresolved_conflict_count']} conflict · {timeline['missing_time_observation_count']} untimed note", "The system showed the conflict and left the untimed note out of confirmed order."],
        ["Action checklist", f"{actions['complete_required_fields']}/{actions['possible_required_fields']} fields", "One of four action owners is intentionally blank."],
        ["Rule request time", f"Median {latency['p50']} ms · 95th percentile {latency['p95']} ms", "Twelve local test calls; no AI model was loaded."],
        ["Test-process memory", f"{BENCHMARK['peak_process_rss_mb']:.1f} MB", "Memory used by the app and test harness on one Mac; excludes model use."],
    ]
    story.append(pdf_table(rows[0], rows[1:], [1.25*inch, 2.05*inch, 3.4*inch], styles))
    story += [
        Paragraph("Edge cases checked", styles["Section"]),
        Paragraph("A note without a time stays unplaced. The 09:10 and 09:11 reports remain a visible disagreement. ‘Generator lagged a bit’ is too vague to become a finding. The low-confidence transcript mistake stays uncertain; the response-time proposal uses the separate event log.", styles["BodyDL"]),
        Paragraph("What was not measured", styles["Section"]),
        Paragraph("No human evaluators took part. The project did not measure report-writing time, reviewer edits, evaluator agreement, cash savings, performance on other scenarios, PDF image-reading (OCR), speech-to-text word error rate, or the quality and speed of a local language model. Word error rate means the share of spoken words that a transcript gets wrong. Optional speech-to-text and local language-model paths were not run.", styles["CalloutDL"]),
        Paragraph("Comparison of approaches", styles["Section"]),
    ]
    story.append(pdf_table(["Approach", "Status", "What that means"], [
        ["Manual evaluator review", "Not measured", "No participant, timed baseline, or record of corrections."],
        ["Fixed local rules", "Measured on one fixture", "Matched its labels. This does not predict results on new exercises."],
        ["Optional local writing model", "Not run", "No quality or speed comparison is available."],
        ["Optional speech-to-text", "Not run", "No real audio, word-error rate, or resource comparison is available."],
    ], [1.55*inch, 1.35*inch, 3.8*inch], styles))
    story += [
        Paragraph("Next evaluation steps", styles["Section"]),
    ]
    story.append(pdf_table(["Step", "What to measure"], [
        ["Work with evaluators", "Time evidence lookup, drafting, edits, and report completion on matched exercises."],
        ["Use an unfamiliar exercise", "Have evaluators label records before analysis; measure missed and incorrect findings, unsupported claims, and timeline agreement."],
        ["Compare workflow and AI assistance", "Compare manual review with deterministic rules on the same records; test optional local AI only after it is configured, against the same labels and reviewer workload."],
        ["Check citations", "Ask independent reviewers whether each cited record actually supports each claim; compare with one selected alternative when partner access permits."],
        ["Test speech-to-text", "With consented English audio, measure transcript errors, time corrections, and computer resources."],
        ["Follow actions", "Measure owner completeness, closure evidence quality, and independently verified completion."],
    ], [1.55*inch, 5.15*inch], styles))
    story.append(Paragraph("Interpretation", styles["Section"]))
    story.append(Paragraph("The demo shows that a source-linked review flow can be built with visible uncertainty and human decision gates. It does not show that AI can outperform evaluators or generalize to unfamiliar exercises. A proposed pilot would compare the current process with DrillLoop on the same permitted records, use evaluator-labeled evidence from an unfamiliar exercise, compare missed and incorrect findings, and agree success thresholds with the partner before measurement.", styles["BodyDL"]))
    doc = SimpleDocTemplate(str(ROOT / "DrillLoop_Evaluation_Report.pdf"), pagesize=letter, rightMargin=.62*inch, leftMargin=.62*inch, topMargin=.62*inch, bottomMargin=.65*inch, title="DrillLoop Evaluation Report")
    doc.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)


def main() -> None:
    build_idea_docx()
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run(
            [soffice, "--headless", "-env:UserInstallation=file:///tmp/drillloop-docs-profile",
             "--convert-to", "pdf", "--outdir", str(ROOT), str(ROOT / "DrillLoop_Idea_Document.docx")],
            check=True, capture_output=True, text=True,
        )
    else:
        build_idea_pdf()
    build_business_pdf()
    build_evaluation_pdf()
    print("Generated DrillLoop_Idea_Document.docx/.pdf, DrillLoop_Business_Case.pdf, and DrillLoop_Evaluation_Report.pdf")


if __name__ == "__main__":
    main()
